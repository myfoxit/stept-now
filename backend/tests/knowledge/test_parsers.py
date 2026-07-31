"""Parser unit tests: pdf, docx, html, markdown, csv, txt, unsupported."""

from __future__ import annotations

import io

import pytest

from app.rag.parsers import ParseError, extract


def build_pdf(text: str) -> bytes:
    """Hand-rolled minimal single-page PDF with a correct xref table."""
    stream = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode()
    objects = [
        b"1 0 obj << /Type /Catalog /Pages 2 0 R >> endobj\n",
        b"2 0 obj << /Type /Pages /Kids [3 0 R] /Count 1 >> endobj\n",
        b"3 0 obj << /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
        b"/Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >> endobj\n",
        b"4 0 obj << /Length %d >> stream\n%s\nendstream endobj\n" % (len(stream), stream),
        b"5 0 obj << /Type /Font /Subtype /Type1 /BaseFont /Helvetica >> endobj\n",
    ]
    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for obj in objects:
        offsets.append(len(out))
        out += obj
    xref_pos = len(out)
    out += f"xref\n0 {len(objects) + 1}\n".encode()
    out += b"0000000000 65535 f \n"
    for offset in offsets:
        out += f"{offset:010d} 00000 n \n".encode()
    out += f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\n".encode()
    out += f"startxref\n{xref_pos}\n%%EOF".encode()
    return bytes(out)


def build_pdf_pages(texts: list[str]) -> bytes:
    """Multi-page variant of `build_pdf` (one text line per page)."""
    count = len(texts)
    font_number = 3 + 2 * count
    kids = " ".join(f"{3 + 2 * index} 0 R" for index in range(count))
    objects = [
        b"1 0 obj << /Type /Catalog /Pages 2 0 R >> endobj\n",
        f"2 0 obj << /Type /Pages /Kids [{kids}] /Count {count} >> endobj\n".encode(),
    ]
    for index, text in enumerate(texts):
        stream = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode()
        objects.append(
            f"{3 + 2 * index} 0 obj << /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            f"/Contents {4 + 2 * index} 0 R /Resources << /Font << /F1 {font_number} 0 R >> >> "
            f">> endobj\n".encode()
        )
        objects.append(
            b"%d 0 obj << /Length %d >> stream\n%s\nendstream endobj\n"
            % (4 + 2 * index, len(stream), stream)
        )
    objects.append(
        f"{font_number} 0 obj << /Type /Font /Subtype /Type1 "
        f"/BaseFont /Helvetica >> endobj\n".encode()
    )
    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for obj in objects:
        offsets.append(len(out))
        out += obj
    xref_pos = len(out)
    out += f"xref\n0 {len(objects) + 1}\n".encode()
    out += b"0000000000 65535 f \n"
    for offset in offsets:
        out += f"{offset:010d} 00000 n \n".encode()
    out += f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\n".encode()
    out += f"startxref\n{xref_pos}\n%%EOF".encode()
    return bytes(out)


def test_pdf_extraction():
    parsed = extract("refund-policy.pdf", build_pdf("Refunds are issued within 30 days."))
    assert "Refunds are issued within 30 days." in parsed.text
    assert parsed.title == "refund policy"
    assert parsed.meta["pages"] == 1
    assert parsed.meta["mime"] == "application/pdf"
    assert "failed_pages" not in parsed.meta


def test_pdf_corrupt_raises():
    with pytest.raises(ParseError):
        extract("broken.pdf", b"%PDF-1.4 not really a pdf")


def test_pdf_page_failure_keeps_the_other_pages(monkeypatch):
    from pypdf._page import PageObject

    original = PageObject.extract_text
    seen = {"pages": 0}

    def flaky(self, *args, **kwargs):
        seen["pages"] += 1
        if seen["pages"] == 2:
            raise ValueError("page 2 is broken")
        return original(self, *args, **kwargs)

    monkeypatch.setattr(PageObject, "extract_text", flaky)
    parsed = extract(
        "handbook.pdf", build_pdf_pages(["Page one facts.", "Broken page.", "Page three facts."])
    )
    assert "Page one facts." in parsed.text
    assert "Page three facts." in parsed.text
    assert "Broken page." not in parsed.text
    assert parsed.meta["pages"] == 3
    assert parsed.meta["failed_pages"] == [2]  # 1-based page numbers


def test_docx_extraction_with_headings():
    import docx

    document = docx.Document()
    document.add_heading("Billing guide", level=1)
    document.add_paragraph("We accept all major credit cards.")
    document.add_heading("Refunds", level=2)
    document.add_paragraph("Full refund within 30 days.")
    buffer = io.BytesIO()
    document.save(buffer)

    parsed = extract("billing.docx", buffer.getvalue())
    assert parsed.title == "Billing guide"
    assert "# Billing guide" in parsed.text
    assert "## Refunds" in parsed.text
    assert "Full refund within 30 days." in parsed.text


def test_docx_tables_become_markdown_in_document_order():
    import docx

    document = docx.Document()
    document.add_heading("Plans", level=1)
    document.add_paragraph("Compare the plans below.")
    table = document.add_table(rows=3, cols=2)
    for row_index, cells in enumerate([("Plan", "Price | month"), ("Pro", "49"), ("Team", "99")]):
        for column_index, value in enumerate(cells):
            table.rows[row_index].cells[column_index].text = value
    document.add_paragraph("Prices are per seat.")
    buffer = io.BytesIO()
    document.save(buffer)

    parsed = extract("plans.docx", buffer.getvalue())
    assert parsed.title == "Plans"
    lines = parsed.text.splitlines()
    assert "| Plan | Price \\| month |" in lines  # header row, pipes escaped
    assert "| --- | --- |" in lines
    assert "| Pro | 49 |" in lines
    assert "| Team | 99 |" in lines
    # order: heading → intro paragraph → table → trailing paragraph
    assert lines.index("Compare the plans below.") < lines.index("| Pro | 49 |")
    assert lines.index("| Team | 99 |") < lines.index("Prices are per seat.")


def test_docx_table_rows_are_capped():
    import docx

    document = docx.Document()
    table = document.add_table(rows=121, cols=2)
    table.rows[0].cells[0].text = "user"
    table.rows[0].cells[1].text = "plan"
    for index in range(1, 121):
        table.rows[index].cells[0].text = f"user{index}"
        table.rows[index].cells[1].text = "pro"
    buffer = io.BytesIO()
    document.save(buffer)

    parsed = extract("users.docx", buffer.getvalue())
    assert "| user100 | pro |" in parsed.text
    assert "| user101 | pro |" not in parsed.text  # capped at 100 data rows
    assert "truncated to the first 100 rows of 120" in parsed.text


def test_html_headings_preserved_as_markdown():
    html = b"""
    <html><head><title>Widget install guide</title></head><body>
      <nav>Skip me not required</nav>
      <h1>Install the widget</h1>
      <p>Copy the embed snippet.</p>
      <h2>Verify identity</h2>
      <p>Sign the user id with HMAC.</p>
      <ul><li>Step one</li><li>Step two</li></ul>
      <script>var ignored = true;</script>
    </body></html>
    """
    parsed = extract("guide.html", html, "text/html")
    assert parsed.title == "Widget install guide"
    assert "# Install the widget" in parsed.text
    assert "## Verify identity" in parsed.text
    assert "- Step one" in parsed.text
    assert "ignored" not in parsed.text


def test_markdown_title_from_first_heading():
    parsed = extract("notes.md", b"intro line\n\n# Real Title\n\nBody text here.")
    assert parsed.title == "Real Title"
    assert "Body text here." in parsed.text


def test_txt_title_from_filename():
    parsed = extract("release_notes.txt", b"Version 2 ships tomorrow.")
    assert parsed.title == "release notes"
    assert parsed.text == "Version 2 ships tomorrow."


def test_csv_becomes_markdown_table_capped():
    rows = ["name,plan"] + [f"user{i},pro" for i in range(250)]
    parsed = extract("customers.csv", "\n".join(rows).encode())
    assert parsed.text.startswith("| name | plan |")
    assert "| user0 | pro |" in parsed.text
    assert "| user199 | pro |" in parsed.text
    assert "| user200 | pro |" not in parsed.text  # capped at 200 data rows
    assert parsed.meta["truncated"] is True
    assert parsed.meta["rows"] == 250


def test_unsupported_type_raises():
    with pytest.raises(ParseError):
        extract("binary.bin", b"\x00\x01\x02", "application/octet-stream")
